#!/usr/bin/env python3
"""gsb_slides — reusable GSB (ธนาคารออมสิน) themed .pptx builder.

ทำไมต้องมีไฟล์นี้:
- Canva/Figma AI เรนเดอร์ภาษาไทย "เพี้ยน" (สระ/วรรณยุกต์หลุด) และคุมธีมแบรนด์ไม่ได้เป๊ะ
- python-pptx + ฟอนต์ Tahoma เรนเดอร์ไทยตรง แต่ต้องตั้ง "complex-script font" (a:cs)
  ด้วย ไม่งั้นบางเครื่องจะ fallback ไปฟอนต์อื่นแล้วสระลอย

วิธีใช้ (สั้น ๆ):
    from gsb_slides import Deck
    d = Deck("ชื่อเรื่อง", subtitle="คำโปรย", meta="ข้อมูลเล็ก ๆ")
    d.title()                                   # สไลด์ปก
    d.agenda(["หัวข้อ 1", "หัวข้อ 2", ...])       # วาระ
    d.cards("หัวข้อ", [("ชื่อการ์ด","บรรทัดล่าง"), ...])
    d.bullets("หัวข้อ", [("ตัวหนา","ส่วนอธิบาย"), "ข้อความล้วน"])
    d.compare("AS IS → TO BE", [("ด้าน","ปัจจุบัน","เป้าหมาย"), ...])
    d.scorecard("ความพร้อม", [("ด้าน","สถานะ","หมายเหตุ"), ...])
    d.phases("หัวข้อ", [("Phase 0","ชื่อ","รายละเอียด"), ...])
    d.envflow("หัวข้อ", [("dev","คำอธิบาย"), ...], principles=[...])
    d.table("หัวข้อ", ["คอลัมน์ซ้าย","คอลัมน์ขวา"], [("แถวซ้าย","แถวขวา"), ...])
    d.closing("ขอบคุณครับ", "Q & A")
    d.save("out.pptx")

ดู SKILL.md สำหรับรายละเอียด/เช็กลิสต์เนื้อหา
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

# ---- GSB / ออมสิน palette ----
PINK    = RGBColor(0xEC, 0x00, 0x8C)   # magenta-pink หลักของแบรนด์
PINKD   = RGBColor(0xB3, 0x00, 0x6A)   # เข้มขึ้น (เส้นใต้ header/แถบ)
PINKLT  = RGBColor(0xEC, 0x9C, 0xC8)   # อ่อน (ตัวรอง)
DARK    = RGBColor(0x22, 0x2B, 0x36)
GREY    = RGBColor(0x5A, 0x66, 0x72)
LIGHT   = RGBColor(0xF4, 0xF1, 0xF6)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
UP      = RGBColor(0x2E, 0x7D, 0x46)
UPBG    = RGBColor(0xDD, 0xEF, 0xE2)
DOWN    = RGBColor(0xB1, 0x3A, 0x3F)
DOWNBG  = RGBColor(0xF6, 0xE0, 0xE1)
AMBER   = RGBColor(0xB9, 0x86, 0x12)
AMBERBG = RGBColor(0xF7, 0xEC, 0xC8)
BLUE    = RGBColor(0x23, 0x4D, 0x80)
BLUEBG  = RGBColor(0xD8, 0xE4, 0xF3)
ACCENTS = [PINK, BLUE, AMBER, UP]      # ใช้วนสีการ์ด/เฟส

# ฟอนต์ที่ปลอดภัยกับภาษาไทยข้ามเครื่อง (มี glyph ไทยครบ)
FONT = "Tahoma"


def _set_font(run, name=FONT):
    """ตั้งฟอนต์ทั้ง latin + complex-script(cs) + east-asian(ea).

    จุดสำคัญ: python-pptx ตั้งให้แค่ latin โดย default พอเจอสระ/วรรณยุกต์ไทย
    PowerPoint จะถือเป็น complex-script แล้ว fallback ไปฟอนต์ธีม ทำให้ 'เพี้ยน'
    ต้อง set a:cs ด้วยถึงจะคุมได้จริง
    """
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:cs", "a:ea"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


class Deck:
    def __init__(self, title, subtitle="", meta="", brand="ธนาคารออมสิน · GSB"):
        self.title_text = title
        self.subtitle = subtitle
        self.meta = meta
        self.brand = brand
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.SW, self.SH = self.prs.slide_width, self.prs.slide_height
        self.BLANK = self.prs.slide_layouts[6]
        self._n = 0

    # ---------- low-level ----------
    def _slide(self, dark=False):
        s = self.prs.slides.add_slide(self.BLANK)
        bg = s.shapes.add_shape(1, 0, 0, self.SW, self.SH)
        bg.fill.solid(); bg.fill.fore_color.rgb = DARK if dark else WHITE
        bg.line.fill.background(); bg.shadow.inherit = False
        s.shapes._spTree.remove(bg._element); s.shapes._spTree.insert(2, bg._element)
        return s

    def _box(self, s, x, y, w, h, fill=None, line=None, line_w=1.0, rounded=False):
        shp = s.shapes.add_shape(5 if rounded else 1,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
        if fill is None:
            shp.fill.background()
        else:
            shp.fill.solid(); shp.fill.fore_color.rgb = fill
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = line; shp.line.width = Pt(line_w)
        shp.shadow.inherit = False
        return shp

    def _text(self, s, x, y, w, h, runs, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP, space_after=4, line_spacing=1.0):
        """runs: list ของ paragraph; แต่ละ paragraph = tuple เดียว หรือ list ของ
        tuple (txt, size, color, bold)."""
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        for i, para in enumerate(runs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align; p.space_after = Pt(space_after)
            p.space_before = Pt(0); p.line_spacing = line_spacing
            if isinstance(para, tuple):
                para = [para]
            for (txt, size, color, bold) in para:
                r = p.add_run(); r.text = txt
                r.font.size = Pt(size); r.font.bold = bold
                r.font.color.rgb = color
                _set_font(r)          # <-- แก้ไทยเพี้ยนตรงนี้
        return tb

    def _header(self, s, kicker, title):
        self._box(s, 0, 0, 13.333, 1.15, fill=PINK)
        self._box(s, 0, 1.15, 13.333, 0.07, fill=PINKD)
        self._text(s, 0.6, 0.16, 12, 0.45, [(kicker, 12, RGBColor(0xFF,0xD6,0xEC), True)])
        self._text(s, 0.6, 0.46, 12, 0.6, [(title, 25, WHITE, True)])

    def _footer(self, s):
        self._text(s, 0.6, 7.08, 9, 0.3, [(self.title_text, 9, GREY, False)])
        self._text(s, 12.0, 7.08, 0.9, 0.3, [(str(self._n), 9, GREY, False)],
                   align=PP_ALIGN.RIGHT)

    # ---------- slide types ----------
    def title(self):
        self._n += 1
        s = self._slide(dark=True)
        self._box(s, 0, 5.0, 13.333, 0.12, fill=PINK)
        self._text(s, 0.9, 1.7, 11.5, 0.5, [(self.brand, 15, RGBColor(0xFF,0xA8,0xD6), True)])
        self._text(s, 0.9, 2.3, 11.5, 1.6,
                   [(self.title_text, 54, WHITE, True),
                    (self.subtitle, 28, PINKLT, True)], line_spacing=1.0)
        if self.meta:
            self._text(s, 0.9, 5.9, 11.5, 0.5, [(self.meta, 13, RGBColor(0xB8,0xC2,0xCE), False)])
        return s

    def section(self, num, title, sub=""):
        """สไลด์คั่นบท."""
        self._n += 1
        s = self._slide(dark=True)
        self._box(s, 0.9, 2.6, 0.16, 1.5, fill=PINK)
        self._text(s, 1.35, 2.55, 11, 0.6, [(f"ส่วนที่ {num}", 16, PINKLT, True)])
        self._text(s, 1.35, 3.05, 11, 0.9, [(title, 40, WHITE, True)])
        if sub:
            self._text(s, 1.35, 4.0, 11, 0.5, [(sub, 15, RGBColor(0xB8,0xC2,0xCE), False)])
        return s

    def agenda(self, items, kicker="วาระการประชุม", title="Agenda"):
        self._n += 1
        s = self._slide(); self._header(s, kicker, title)
        runs = []
        for it in items:
            if isinstance(it, tuple):
                lead, rest = it
                runs.append([("•  ", 17, PINK, True), (lead, 17, DARK, True),
                             (rest, 17, GREY, False)])
            else:
                runs.append([("•  ", 17, PINK, True), (it, 17, DARK, False)])
        self._text(s, 1.1, 1.6, 11, 5.3, runs, space_after=10, line_spacing=1.05)
        self._footer(s); return s

    def bullets(self, title, items, kicker="", size=15, intro=None):
        self._n += 1
        s = self._slide(); self._header(s, kicker, title)
        y = 1.5
        if intro:
            self._text(s, 0.6, y, 12.1, 0.7, [(intro, 14, GREY, False)]); y = 2.15
        runs = []
        for it in items:
            if isinstance(it, tuple):
                lead, rest = it
                runs.append([("•  ", size, PINK, True), (lead, size, DARK, True),
                             (rest, size, GREY, False)])
            else:
                runs.append([("•  ", size, PINK, True), (it, size, DARK, False)])
        self._text(s, 0.9, y, 11.7, 5, runs, space_after=10, line_spacing=1.1)
        self._footer(s); return s

    def cards(self, title, cards, kicker="", cols=3, intro=None):
        """cards: list ของ (t1, t2) หรือ (t1, t2, t3)."""
        self._n += 1
        s = self._slide(); self._header(s, kicker, title)
        y0 = 1.5
        if intro:
            self._text(s, 0.6, 1.4, 12.1, 0.7, [(intro, 14, GREY, False)]); y0 = 2.35
        gx, gy = 0.18, 0.2
        cw = (12.13 - gx * (cols - 1)) / cols
        ch = 1.5 if cols >= 3 else 1.2
        for i, c in enumerate(cards):
            t1 = c[0]; t2 = c[1] if len(c) > 1 else ""; t3 = c[2] if len(c) > 2 else ""
            col = ACCENTS[i % len(ACCENTS)]
            cx = 0.6 + (i % cols) * (cw + gx)
            cy = y0 + (i // cols) * (ch + gy)
            self._box(s, cx, cy, cw, ch, fill=LIGHT, line=RGBColor(0xD9,0xCE,0xD6), rounded=True)
            self._box(s, cx, cy, 0.13, ch, fill=col)
            self._text(s, cx+0.28, cy+0.16, cw-0.4, 0.4, [(t1, 16, DARK, True)])
            if t2:
                self._text(s, cx+0.28, cy+0.62, cw-0.4, 0.35, [(t2, 12.5, PINKD, True)])
            if t3:
                self._text(s, cx+0.28, cy+1.0, cw-0.4, 0.4, [(t3, 12, GREY, False)],
                           line_spacing=1.05)
        self._footer(s); return s

    def phases(self, title, phases, kicker="แผนงาน", intro=None):
        """phases: list ของ (tag, name, detail)."""
        self._n += 1
        s = self._slide(); self._header(s, kicker, title)
        y = 1.55
        if intro:
            self._text(s, 0.6, 1.4, 12, 0.4, [(intro, 14, GREY, False)]); y = 1.95
        n = len(phases)
        bh = min(1.05, (6.8 - y) / max(1, n) - 0.15)
        for i, ph in enumerate(phases):
            tag, name, detail = ph[0], ph[1], ph[2]
            col = ACCENTS[i % len(ACCENTS)]
            self._box(s, 0.6, y, 0.13, bh, fill=col)
            self._box(s, 0.9, y, 1.7, bh, fill=col, rounded=True)
            self._text(s, 0.9, y, 1.7, bh, [(tag, 16, WHITE, True)],
                       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            self._box(s, 2.75, y, 9.98, bh, fill=LIGHT, line=RGBColor(0xDD,0xD3,0xDA), rounded=True)
            self._text(s, 3.05, y+0.1, 9.5, 0.42, [(name, 16, DARK, True)])
            self._text(s, 3.05, y+0.52, 9.5, 0.45, [(detail, 12, GREY, False)], line_spacing=1.05)
            y += bh + 0.15
        self._footer(s); return s

    def envflow(self, title, envs, principles=None, kicker="การเลื่อนระดับระบบ"):
        """envs: list ของ (name, desc). principles: list ของ (bold, rest) หรือ str."""
        self._n += 1
        s = self._slide(); self._header(s, kicker, title)
        cw = 2.85; gap = 0.45; x = 0.6; y = 1.7
        for i, (e, d) in enumerate(envs):
            col = ACCENTS[i % len(ACCENTS)]
            self._box(s, x, y, cw, 1.5, fill=col, rounded=True)
            self._text(s, x, y+0.22, cw, 0.6, [(e.upper(), 26, WHITE, True)], align=PP_ALIGN.CENTER)
            self._text(s, x+0.15, y+0.85, cw-0.3, 0.55, [(d, 12, WHITE, False)], align=PP_ALIGN.CENTER)
            if i < len(envs) - 1:
                self._text(s, x+cw-0.05, y+0.45, gap+0.1, 0.6, [("→", 30, GREY, True)],
                           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            x += cw + gap
        if principles:
            self._text(s, 0.6, 3.55, 12, 0.4, [("หลักการสำคัญ", 15, PINKD, True)])
            runs = []
            for p in principles:
                if isinstance(p, tuple):
                    lead, rest = p
                    runs.append([("•  ", 14, PINK, True), (lead, 14, DARK, True), (rest, 14, GREY, False)])
                else:
                    runs.append([("•  ", 14, PINK, True), (p, 14, DARK, False)])
            self._text(s, 0.9, 4.05, 11.7, 3, runs, space_after=8, line_spacing=1.1)
        self._footer(s); return s

    def table(self, title, headers, rows, kicker="", widths=None, sizes=None):
        """headers: [ซ้าย, ขวา] หรือมากกว่า. rows: list ของ tuple ตามจำนวน header."""
        self._n += 1
        s = self._slide(); self._header(s, kicker, title)
        ncol = len(headers)
        if widths is None:
            widths = [12.13 / ncol] * ncol
        total = sum(widths); widths = [w * 12.13 / total for w in widths]
        if sizes is None:
            sizes = [12.5] * ncol
        y = 1.55
        x = 0.6
        for i, h in enumerate(headers):
            self._box(s, x, y, widths[i] - 0.05, 0.45, fill=PINK if i == 0 else PINKD)
            self._text(s, x+0.15, y+0.05, widths[i]-0.3, 0.4, [(h, 13, WHITE, True)])
            x += widths[i]
        y += 0.45
        rh = min(0.92, (6.75 - y) / max(1, len(rows)))
        for i, row in enumerate(rows):
            bg = WHITE if i % 2 == 0 else LIGHT
            x = 0.6
            for j, cell in enumerate(row):
                self._box(s, x, y, widths[j] - 0.05, rh, fill=bg, line=RGBColor(0xE3,0xDB,0xE0))
                col = DARK if j == 0 else GREY
                self._text(s, x+0.15, y+0.03, widths[j]-0.3, rh, [(cell, sizes[j], col, j == 0)],
                           anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.02)
                x += widths[j]
            y += rh
        self._footer(s); return s

    def compare(self, title, rows, kicker="เทียบสถานะ",
                left="AS IS — ปัจจุบัน", right="TO BE — เป้าหมาย"):
        """ตาราง 3 คอลัมน์ (ด้าน / AS IS / TO BE). rows: list ของ (ด้าน, as_is, to_be)."""
        self._n += 1
        s = self._slide(); self._header(s, kicker, title)
        cx1, cw1 = 0.6, 2.4
        cx2, cw2 = 3.1, 4.7
        cx3, cw3 = 7.9, 4.83
        y = 1.5
        self._box(s, cx1, y, cw1, 0.5, fill=PINKD)
        self._box(s, cx2, y, cw2, 0.5, fill=DOWN)
        self._box(s, cx3, y, cw3, 0.5, fill=UP)
        self._text(s, cx1+0.15, y+0.08, cw1-0.3, 0.4, [("ด้าน", 13, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)
        self._text(s, cx2+0.15, y+0.08, cw2-0.3, 0.4, [(left, 13, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)
        self._text(s, cx3+0.15, y+0.08, cw3-0.3, 0.4, [(right, 13, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)
        y += 0.5
        rh = min(0.63, (6.7 - y) / max(1, len(rows)))
        for i, (dan, a, b) in enumerate(rows):
            bg = WHITE if i % 2 == 0 else LIGHT
            for cx, cw in ((cx1, cw1), (cx2, cw2), (cx3, cw3)):
                self._box(s, cx, y, cw, rh, fill=bg, line=RGBColor(0xE3,0xDB,0xE0))
            self._text(s, cx1+0.15, y+0.03, cw1-0.28, rh, [(dan, 12, DARK, True)], anchor=MSO_ANCHOR.MIDDLE)
            self._text(s, cx2+0.15, y+0.03, cw2-0.28, rh, [(a, 11.5, GREY, False)], anchor=MSO_ANCHOR.MIDDLE)
            self._text(s, cx3+0.15, y+0.03, cw3-0.28, rh, [(b, 11.5, DARK, False)], anchor=MSO_ANCHOR.MIDDLE)
            y += rh
        self._footer(s); return s

    def scorecard(self, title, rows, kicker="ประเมินความพร้อม", intro=None,
                  legend=True):
        """rows: list ของ (ด้าน, ระดับ, หมายเหตุ) โดย ระดับ = 'ready'|'partial'|'gap'."""
        self._n += 1
        s = self._slide(); self._header(s, kicker, title)
        y = 1.45
        if intro:
            self._text(s, 0.6, y, 12.1, 0.4, [(intro, 13, GREY, False)]); y = 1.9
        LABEL = {"ready": ("พร้อม", UP, UPBG),
                 "partial": ("บางส่วน", AMBER, AMBERBG),
                 "gap": ("ยังไม่มี", DOWN, DOWNBG)}
        cx1, cw1 = 0.6, 3.3
        cx2, cw2 = 4.0, 1.5
        cx3, cw3 = 5.65, 7.08
        self._box(s, cx1, y, cw1, 0.42, fill=PINKD)
        self._box(s, cx2, y, cw2, 0.42, fill=PINKD)
        self._box(s, cx3, y, cw3, 0.42, fill=PINKD)
        self._text(s, cx1+0.15, y+0.04, cw1-0.3, 0.36, [("ด้าน", 12.5, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)
        self._text(s, cx2+0.15, y+0.04, cw2-0.3, 0.36, [("สถานะ", 12.5, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)
        self._text(s, cx3+0.15, y+0.04, cw3-0.3, 0.36, [("สิ่งที่ต้องทำเพื่อขึ้น Production", 12.5, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)
        y += 0.42
        bottom = 6.55 if legend else 6.9
        rh = min(0.6, (bottom - y) / max(1, len(rows)))
        for i, (dan, lvl, note) in enumerate(rows):
            bg = WHITE if i % 2 == 0 else LIGHT
            txt, fg, chipbg = LABEL.get(lvl, LABEL["gap"])
            for cx, cw in ((cx1, cw1), (cx2, cw2), (cx3, cw3)):
                self._box(s, cx, y, cw, rh, fill=bg, line=RGBColor(0xE3,0xDB,0xE0))
            self._text(s, cx1+0.15, y+0.02, cw1-0.3, rh, [(dan, 12, DARK, True)], anchor=MSO_ANCHOR.MIDDLE)
            ch = min(0.3, rh - 0.12)
            self._box(s, cx2+0.18, y+(rh-ch)/2, cw2-0.36, ch, fill=chipbg, line=fg, rounded=True)
            self._text(s, cx2+0.18, y+(rh-ch)/2, cw2-0.36, ch, [(txt, 10.5, fg, True)],
                       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            self._text(s, cx3+0.15, y+0.02, cw3-0.3, rh, [(note, 11, GREY, False)],
                       anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.02)
            y += rh
        if legend:
            lx = 0.6
            for key in ("ready", "partial", "gap"):
                txt, fg, chipbg = LABEL[key]
                self._box(s, lx, 6.72, 1.35, 0.3, fill=chipbg, line=fg, rounded=True)
                self._text(s, lx, 6.72, 1.35, 0.3, [(txt, 10.5, fg, True)],
                           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                lx += 1.5
        self._footer(s); return s

    def kpis(self, title, items, kicker="ตัวเลขสำคัญ", intro=None):
        """items: list ของ (ตัวเลข, ป้าย, คำอธิบาย)."""
        self._n += 1
        s = self._slide(); self._header(s, kicker, title)
        y = 1.6
        if intro:
            self._text(s, 0.6, 1.42, 12.1, 0.4, [(intro, 13, GREY, False)]); y = 2.0
        n = len(items)
        gx = 0.2
        cw = (12.13 - gx * (n - 1)) / n
        for i, (big, label, note) in enumerate(items):
            col = ACCENTS[i % len(ACCENTS)]
            cx = 0.6 + i * (cw + gx)
            self._box(s, cx, y, cw, 2.1, fill=LIGHT, line=RGBColor(0xD9,0xCE,0xD6), rounded=True)
            self._box(s, cx, y, cw, 0.12, fill=col)
            self._text(s, cx, y+0.35, cw, 0.9, [(big, 40, col, True)], align=PP_ALIGN.CENTER)
            self._text(s, cx+0.15, y+1.25, cw-0.3, 0.4, [(label, 14, DARK, True)], align=PP_ALIGN.CENTER)
            self._text(s, cx+0.15, y+1.65, cw-0.3, 0.4, [(note, 11, GREY, False)],
                       align=PP_ALIGN.CENTER, line_spacing=1.05)
        self._footer(s); return s

    def closing(self, big="ขอบคุณครับ", small="Q & A"):
        self._n += 1
        s = self._slide(dark=True)
        self._box(s, 0, 4.7, 13.333, 0.12, fill=PINK)
        self._text(s, 0.9, 2.6, 11.5, 1.0, [(big, 46, WHITE, True), (small, 24, PINKLT, True)])
        if self.meta:
            self._text(s, 0.9, 5.0, 11.5, 0.5, [(self.meta, 13, RGBColor(0xB8,0xC2,0xCE), False)])
        return s

    def save(self, out):
        self.prs.save(out)
        return out, len(self.prs.slides._sldIdLst)


if __name__ == "__main__":
    # เดโมเล็ก ๆ พิสูจน์ว่าไทยไม่เพี้ยน + ธีม GSB
    d = Deck("เดโม GSB Slides", subtitle="ทดสอบภาษาไทย สระ วรรณยุกต์",
             meta="ธนาคารออมสิน · ตัวอย่างจาก skill gsb-slides")
    d.title()
    d.bullets("ทดสอบการเรนเดอร์ภาษาไทย", [
        ("สระบน/ล่าง", " — ก็ ปิ่ ปื้ ณ์ ญ์ ฐ์ ครบไหม"),
        ("วรรณยุกต์", " — น้ำ เก้า ผู้ ที่ สู้ ๆ"),
        "ประโยคยาว: ธนาคารออมสินพัฒนาระบบงานอย่างต่อเนื่อง",
    ])
    d.scorecard("ตัวอย่าง Scorecard", [
        ("ความปลอดภัย", "gap", "ยังไม่มี secret store / audit log"),
        ("การสำรองข้อมูล", "partial", "มี backup แต่ยังไม่เคยทดสอบ restore"),
        ("โครงสร้างพื้นฐาน", "ready", "docker stack ครบแล้ว"),
    ])
    d.closing()
    path, n = d.save("gsb-slides-demo.pptx")
    print("wrote", path, "·", n, "slides")
